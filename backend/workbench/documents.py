"""Document conversion, immutable versions and reviewable edits."""

from __future__ import annotations
import copy
import io
import json
import zipfile
from fastapi import HTTPException
from workbench.store import BlobStore, Version, now

MAX_UPLOAD = 32 * 1024 * 1024
MAX_TEXT = 2_000_000
NODE_TYPES = {
    "doc",
    "paragraph",
    "heading",
    "text",
    "bulletList",
    "orderedList",
    "listItem",
    "blockquote",
    "hardBreak",
    "horizontalRule",
    "codeBlock",
    "table",
    "tableRow",
    "tableCell",
    "tableHeader",
    "image",
    "pageBreak",
}


def text_of(node):
    if isinstance(node, str):
        return node
    if not isinstance(node, dict):
        return ""
    if node.get("type") == "text":
        return node.get("text", "")
    return (
        "\n"
        if node.get("type") in {"doc", "bulletList", "orderedList", "table", "tableRow"}
        else ""
    ).join(text_of(x) for x in node.get("content", []))


def native(text=""):
    return {
        "type": "doc",
        "content": [
            {"type": "paragraph", "content": [{"type": "text", "text": line}]}
            if line
            else {"type": "paragraph"}
            for line in text.split("\n")
        ],
    }


def validate_content(content):
    if (
        not isinstance(content, dict)
        or content.get("type") != "doc"
        or len(json.dumps(content)) > MAX_TEXT * 3
    ):
        raise HTTPException(422, "无效或过大的文档")
    count = 0

    def visit(node, depth=0):
        nonlocal count
        count += 1
        if (
            depth > 30
            or count > 50000
            or not isinstance(node, dict)
            or node.get("type") not in NODE_TYPES
        ):
            raise HTTPException(422, "不支持的文档结构")
        if "text" in node and not isinstance(node["text"], str):
            raise HTTPException(422, "无效文本")
        for mark in node.get("marks", []):
            if mark.get("type") not in {
                "bold",
                "italic",
                "strike",
                "underline",
                "highlight",
                "code",
                "link",
            }:
                raise HTTPException(422, "不支持的文字格式")
            if mark.get("type") == "link" and not str(
                mark.get("attrs", {}).get("href", "")
            ).startswith(("https://", "http://", "mailto:")):
                raise HTTPException(422, "无效链接")
        if node.get("type") == "image":
            src = str(node.get("attrs", {}).get("src", ""))
            if not src.startswith("data:image/") or not src.split(";", 1)[0] in {
                "data:image/png",
                "data:image/jpeg",
                "data:image/webp",
            }:
                raise HTTPException(422, "图片必须为内嵌 PNG/JPEG/WebP")
        for child in node.get("content", []):
            visit(child, depth + 1)

    visit(content)
    return content


def extract(content: bytes, filename: str):
    suffix = filename.rsplit(".", 1)[-1].lower()
    if suffix == "pdf":
        import fitz

        try:
            with fitz.open(stream=content, filetype="pdf") as pdf:
                if pdf.is_encrypted or len(pdf) > 1000:
                    raise ValueError()
                text = "\n".join(page.get_text() for page in pdf)[:MAX_TEXT]
                return {"format": "pdf", "pages": len(pdf), "text": text}
        except Exception:
            raise HTTPException(422, "PDF 无法读取或超出处理范围")
    if suffix == "docx":
        from docx import Document

        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if sum(x.file_size for x in archive.infolist()) > MAX_UPLOAD * 5:
                    raise ValueError()
            doc = Document(io.BytesIO(content))
            blocks = []
            from docx.oxml.ns import qn
            from docx.text.paragraph import Paragraph
            from docx.table import Table

            for child in doc.element.body:
                if child.tag == qn("w:p"):
                    p = Paragraph(child, doc)
                    runs = []
                    for run in p.runs:
                        if not run.text:
                            continue
                        marks = [
                            {"type": name}
                            for attr, name in [
                                ("bold", "bold"),
                                ("italic", "italic"),
                                ("underline", "underline"),
                            ]
                            if getattr(run, attr)
                        ]
                        if run.font.strike:
                            marks.append({"type": "strike"})
                        if run.font.highlight_color is not None:
                            marks.append({"type": "highlight"})
                        runs.append({"type": "text", "text": run.text, "marks": marks})
                    block = {"type": "paragraph", "content": runs}
                    if p.style and p.style.name.startswith("Heading"):
                        block.update(
                            type="heading",
                            attrs={"level": min(3, int(p.style.name[-1]))},
                        )
                    blocks.append(block)
                elif child.tag == qn("w:tbl"):
                    table = Table(child, doc)
                    blocks.append(
                        {
                            "type": "table",
                            "content": [
                                {
                                    "type": "tableRow",
                                    "content": [
                                        {
                                            "type": "tableCell",
                                            "content": native(cell.text)["content"],
                                        }
                                        for cell in row.cells
                                    ],
                                }
                                for row in table.rows
                            ],
                        }
                    )
            value = {"type": "doc", "content": blocks or native()["content"]}
            return {
                "format": "docx",
                "text": text_of(value)[:MAX_TEXT],
                "import_content": value,
            }
        except HTTPException:
            raise
        except Exception:
            raise HTTPException(422, "Word 文件无法读取")
    if suffix in {"txt", "md"}:
        try:
            text = content.decode("utf-8-sig")[:MAX_TEXT]
        except UnicodeDecodeError:
            raise HTTPException(422, "文本文件需要 UTF-8 编码")
        return {"format": suffix, "text": text, "import_content": native(text)}
    if suffix in {"png", "jpg", "jpeg", "webp"}:
        import fitz

        try:
            with fitz.open(
                stream=content, filetype=suffix if suffix != "jpg" else "jpeg"
            ) as image:
                if image[0].rect.width * image[0].rect.height > 40_000_000:
                    raise ValueError()
        except Exception:
            raise HTTPException(422, "图片无法读取或像素过大")
        return {"format": "image", "text": ""}
    raise HTTPException(422, "支持 PDF、DOCX、TXT、Markdown 和图片")


def save_version(session, item):
    session.add(
        Version(
            document_id=item.id,
            owner=item.owner,
            revision=item.revision,
            data=copy.deepcopy(item.data),
        )
    )


def update_content(session, item, content):
    validate_content(content)
    item.data = {**item.data, "content": content, "format": "native"}
    item.searchable = text_of(content)
    item.revision += 1
    item.updated_at = now()
    save_version(session, item)


def replace_text(content, before, after):
    """Apply only an unambiguous exact range, preserving unaffected nodes and marks."""
    result = copy.deepcopy(content)
    if not before or text_of(result).count(before) != 1:
        raise HTTPException(409, "原文已变化或出现多次，请重新选择片段")

    def visit(node):
        if node.get("type") in {"paragraph", "heading", "codeBlock"}:
            leaves = []

            def collect(n):
                if n.get("type") == "text":
                    leaves.append(n)
                for c in n.get("content", []):
                    collect(c)

            collect(node)
            joined = "".join(x.get("text", "") for x in leaves)
            start = joined.find(before)
            if start >= 0:
                end = start + len(before)
                offset = 0
                inserted = False
                for leaf in leaves:
                    original = leaf["text"]
                    a, b = offset, offset + len(original)
                    if b > start and a < end:
                        leaf["text"] = (
                            original[: max(0, start - a)]
                            + (after if not inserted else "")
                            + original[max(0, end - a) :]
                        )
                        inserted = True
                    offset = b
                return True
        for child in node.get("content", []):
            if visit(child):
                return True
        return False

    if not visit(result):
        raise HTTPException(409, "请按单个段落提交修改建议；跨段落选区可以提问")

    def prune(node):
        if "content" in node:
            node["content"] = [
                prune(x)
                for x in node["content"]
                if x.get("type") != "text" or x.get("text")
            ]
        return node

    return prune(result)


def export_docx(content):
    from docx import Document
    from docx.shared import Pt
    from docx.oxml.ns import qn
    import base64

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "宋体"
    style.font.size = Pt(12)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    def append(node, host=doc):
        kind = node.get("type")
        if kind == "table":
            rows = node.get("content", [])
            if not rows:
                return
            table = host.add_table(
                rows=len(rows), cols=max(len(r.get("content", [])) for r in rows)
            )
            table.style = "Table Grid"
            for i, row in enumerate(rows):
                for j, cell in enumerate(row.get("content", [])):
                    table.cell(i, j).text = text_of(cell)
        elif kind == "pageBreak":
            host.add_page_break()
        elif kind == "image":
            from docx.shared import Inches

            raw = node.get("attrs", {}).get("src", "").split(",", 1)[-1]
            host.add_picture(io.BytesIO(base64.b64decode(raw)), width=Inches(5))
        elif kind in {"paragraph", "heading", "codeBlock"}:
            p = (
                host.add_heading("", level=node.get("attrs", {}).get("level", 1))
                if kind == "heading"
                else host.add_paragraph()
            )
            for child in node.get("content", []):
                run = p.add_run(
                    "\n" if child.get("type") == "hardBreak" else child.get("text", "")
                )
                marks = {x["type"] for x in child.get("marks", [])}
                run.bold, run.italic, run.underline = (
                    "bold" in marks,
                    "italic" in marks,
                    "underline" in marks,
                )
                run.font.strike = "strike" in marks
                if "highlight" in marks:
                    from docx.enum.text import WD_COLOR_INDEX

                    run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        else:
            for child in node.get("content", []):
                append(child, host)

    append(content)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
