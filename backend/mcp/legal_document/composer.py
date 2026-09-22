"""
模块描述：Block-based DOCX 文书编排器。
LLM 提交结构化 blocks 数组，本模块将其渲染为符合中文法律文书排版规范的 docx 文件。
"""

from __future__ import annotations

from typing import Any

from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Cm

from .styles import (
    apply_body_style,
    apply_heading_style,
    apply_signature_style,
    apply_table_cell_style,
    apply_title_style,
    configure_page,
)
from .text_cleaner import clean_block_text, merge_refs


SUPPORTED_BLOCK_TYPES = {
    "title",
    "heading",
    "paragraph",
    "ordered_list",
    "unordered_list",
    "table",
    "signature_block",
    "page_break",
    "blank_line",
}


def _add_title(doc, block: dict) -> list[dict]:
    text = str(block.get("text", "")).strip()
    if not text:
        return []
    cleaned, refs = clean_block_text(text)
    if not cleaned:
        return refs
    para = doc.add_paragraph()
    para.add_run(cleaned)
    apply_title_style(para)
    return refs


def _add_heading(doc, block: dict) -> list[dict]:
    text = str(block.get("text", "")).strip()
    if not text:
        return []
    cleaned, refs = clean_block_text(text)
    if not cleaned:
        return refs
    level = int(block.get("level", 1))
    level = max(1, min(3, level))
    para = doc.add_paragraph()
    para.add_run(cleaned)
    apply_heading_style(para, level)
    return refs


def _add_paragraph(doc, block: dict) -> list[dict]:
    text = str(block.get("text", "")).rstrip()
    if not text:
        return []
    cleaned, refs = clean_block_text(text)
    if not cleaned:
        return refs
    indent = block.get("indent", True)
    align = str(block.get("align", "justify"))
    # 内嵌换行符按段落拆分,保持每段独立缩进
    for part in cleaned.split("\n"):
        if not part.strip():
            continue
        para = doc.add_paragraph()
        para.add_run(part.strip())
        apply_body_style(para, indent=bool(indent), align=align)
    return refs


def _add_ordered_list(doc, block: dict) -> list[dict]:
    items = block.get("items") or []
    refs: list[dict] = []
    for i, item in enumerate(items, start=1):
        text = str(item).strip()
        if not text:
            continue
        cleaned, item_refs = clean_block_text(text)
        if not cleaned:
            continue
        refs = merge_refs(refs, item_refs)
        para = doc.add_paragraph()
        para.add_run(f"{i}. {cleaned}")
        apply_body_style(para, indent=True, align="justify")
    return refs


def _add_unordered_list(doc, block: dict) -> list[dict]:
    items = block.get("items") or []
    refs: list[dict] = []
    for item in items:
        text = str(item).strip()
        if not text:
            continue
        cleaned, item_refs = clean_block_text(text)
        if not cleaned:
            continue
        refs = merge_refs(refs, item_refs)
        para = doc.add_paragraph()
        para.add_run(f"· {cleaned}")
        apply_body_style(para, indent=True, align="justify")
    return refs


def _add_table(doc, block: dict) -> list[dict]:
    header = block.get("header") or []
    rows = block.get("rows") or []
    if not header and not rows:
        return []
    col_count = len(header) if header else max((len(r) for r in rows), default=0)
    if col_count == 0:
        return []
    refs: list[dict] = []
    table = doc.add_table(rows=0, cols=col_count)
    table.style = "Table Grid"
    if header:
        cells = table.add_row().cells
        for i in range(col_count):
            raw = str(header[i]) if i < len(header) else ""
            cleaned, cell_refs = clean_block_text(raw)
            refs = merge_refs(refs, cell_refs)
            cells[i].text = cleaned
            apply_table_cell_style(cells[i], header=True)
    for row in rows:
        cells = table.add_row().cells
        for i in range(col_count):
            raw = str(row[i]) if i < len(row) else ""
            cleaned, cell_refs = clean_block_text(raw)
            refs = merge_refs(refs, cell_refs)
            cells[i].text = cleaned
            apply_table_cell_style(cells[i], header=False)
    return refs


def _add_signature_block(doc, block: dict) -> list[dict]:
    """落款块：包含签名 / 单位 / 日期等，每行右对齐独立成段。"""
    lines: list[str] = []
    if block.get("lines"):
        lines = [str(x).strip() for x in block["lines"] if str(x).strip()]
    else:
        if block.get("signer"):
            lines.append(f"具状人：{block['signer']}")
        if block.get("entity"):
            lines.append(str(block["entity"]))
        if block.get("date"):
            lines.append(str(block["date"]))
    refs: list[dict] = []
    for line in lines:
        cleaned, line_refs = clean_block_text(line)
        if not cleaned:
            continue
        refs = merge_refs(refs, line_refs)
        para = doc.add_paragraph()
        para.add_run(cleaned)
        apply_signature_style(para)
    return refs


def _add_page_break(doc, _block: dict) -> list[dict]:
    para = doc.add_paragraph()
    run = para.add_run()
    run.add_break(WD_BREAK.PAGE)
    return []


def _add_blank_line(doc, _block: dict) -> list[dict]:
    para = doc.add_paragraph()
    apply_body_style(para, indent=False, align="left")
    return []


_DISPATCH = {
    "title": _add_title,
    "heading": _add_heading,
    "paragraph": _add_paragraph,
    "ordered_list": _add_ordered_list,
    "unordered_list": _add_unordered_list,
    "table": _add_table,
    "signature_block": _add_signature_block,
    "page_break": _add_page_break,
    "blank_line": _add_blank_line,
}


def compose_docx(blocks: list[dict], output_path: str) -> dict:
    """将 blocks 渲染为 docx 文件。返回渲染统计与从 text 中提取的引用列表。"""
    doc = Document()
    configure_page(doc)

    rendered: dict[str, int] = {}
    skipped_unknown: list[str] = []
    extracted_refs: list[dict] = []

    for block in blocks:
        if not isinstance(block, dict):
            continue
        btype = block.get("type")
        if not isinstance(btype, str):
            skipped_unknown.append(str(btype))
            continue
        handler = _DISPATCH.get(btype)
        if handler is None:
            skipped_unknown.append(btype)
            continue
        block_refs = handler(doc, block) or []
        extracted_refs = merge_refs(extracted_refs, block_refs)
        rendered[btype] = rendered.get(btype, 0) + 1

    doc.save(output_path)

    return {
        "rendered_block_counts": rendered,
        "skipped_unknown_types": skipped_unknown,
        "extracted_refs": extracted_refs,
    }
