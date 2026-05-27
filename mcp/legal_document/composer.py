"""
模块描述：Block-based DOCX 文书编排器。
LLM 提交结构化 blocks 数组，本模块将其渲染为符合中文法律文书排版规范的 docx 文件。
"""

from __future__ import annotations

from typing import Any

from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Cm

from .styles import (
    apply_body_style,
    apply_heading_style,
    apply_signature_style,
    apply_table_cell_style,
    apply_title_style,
    configure_page,
)


SUPPORTED_BLOCK_TYPES = {
    "title",
    "heading",
    "paragraph",
    "ordered_list",
    "unordered_list",
    "table",
    "signature_block",
    "page_break",
    "blank_line",
}


def _add_title(doc, block: dict) -> None:
    text = str(block.get("text", "")).strip()
    if not text:
        return
    para = doc.add_paragraph()
    para.add_run(text)
    apply_title_style(para)


def _add_heading(doc, block: dict) -> None:
    text = str(block.get("text", "")).strip()
    if not text:
        return
    level = int(block.get("level", 1))
    level = max(1, min(3, level))
    para = doc.add_paragraph()
    para.add_run(text)
    apply_heading_style(para, level)


def _add_paragraph(doc, block: dict) -> None:
    text = str(block.get("text", "")).rstrip()
    if not text:
        return
    indent = block.get("indent", True)
    align = str(block.get("align", "justify"))
    # 内嵌换行符按段落拆分,保持每段独立缩进
    parts = text.split("\n")
    for part in parts:
        if not part.strip():
            continue
        para = doc.add_paragraph()
        para.add_run(part.strip())
        apply_body_style(para, indent=bool(indent), align=align)


def _add_ordered_list(doc, block: dict) -> None:
    items = block.get("items") or []
    for i, item in enumerate(items, start=1):
        text = str(item).strip()
        if not text:
            continue
        para = doc.add_paragraph()
        para.add_run(f"{i}. {text}")
        apply_body_style(para, indent=True, align="justify")


def _add_unordered_list(doc, block: dict) -> None:
    items = block.get("items") or []
    for item in items:
        text = str(item).strip()
        if not text:
            continue
        para = doc.add_paragraph()
        para.add_run(f"· {text}")
        apply_body_style(para, indent=True, align="justify")


def _add_table(doc, block: dict) -> None:
    header = block.get("header") or []
    rows = block.get("rows") or []
    if not header and not rows:
        return
    col_count = len(header) if header else max((len(r) for r in rows), default=0)
    if col_count == 0:
        return
    table = doc.add_table(rows=0, cols=col_count)
    table.style = "Table Grid"
    if header:
        cells = table.add_row().cells
        for i in range(col_count):
            cells[i].text = str(header[i]) if i < len(header) else ""
            apply_table_cell_style(cells[i], header=True)
    for row in rows:
        cells = table.add_row().cells
        for i in range(col_count):
            cells[i].text = str(row[i]) if i < len(row) else ""
            apply_table_cell_style(cells[i], header=False)


def _add_signature_block(doc, block: dict) -> None:
    """落款块：包含签名 / 单位 / 日期等，每行右对齐独立成段。"""
    lines: list[str] = []
    if block.get("lines"):
        lines = [str(x).strip() for x in block["lines"] if str(x).strip()]
    else:
        if block.get("signer"):
            lines.append(f"具状人：{block['signer']}")
        if block.get("entity"):
            lines.append(str(block["entity"]))
        if block.get("date"):
            lines.append(str(block["date"]))
    for line in lines:
        para = doc.add_paragraph()
        para.add_run(line)
        apply_signature_style(para)


def _add_page_break(doc, _block: dict) -> None:
    para = doc.add_paragraph()
    run = para.add_run()
    run.add_break(WD_BREAK.PAGE)


def _add_blank_line(doc, _block: dict) -> None:
    para = doc.add_paragraph()
    apply_body_style(para, indent=False, align="left")


_DISPATCH = {
    "title": _add_title,
    "heading": _add_heading,
    "paragraph": _add_paragraph,
    "ordered_list": _add_ordered_list,
    "unordered_list": _add_unordered_list,
    "table": _add_table,
    "signature_block": _add_signature_block,
    "page_break": _add_page_break,
    "blank_line": _add_blank_line,
}


def compose_docx(blocks: list[dict], output_path: str) -> dict:
    """将 blocks 渲染为 docx 文件。返回渲染统计。"""
    doc = Document()
    configure_page(doc)

    rendered: dict[str, int] = {}
    skipped_unknown: list[str] = []

    for block in blocks:
        if not isinstance(block, dict):
            continue
        btype = block.get("type")
        if not isinstance(btype, str):
            skipped_unknown.append(str(btype))
            continue
        handler = _DISPATCH.get(btype)
        if handler is None:
            skipped_unknown.append(btype)
            continue
        handler(doc, block)
        rendered[btype] = rendered.get(btype, 0) + 1

    doc.save(output_path)

    return {
        "rendered_block_counts": rendered,
        "skipped_unknown_types": skipped_unknown,
    }
